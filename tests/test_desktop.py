"""Desktop integration: real engines, persistence, safe output and cancellation."""

from __future__ import annotations

import base64
import io
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from config import ConvertOptions
from desktop.app import DesktopAPI, InstanceLock
from desktop.preview import markdown_html
from desktop.service import Service, validate_options
from desktop.store import Store


@pytest.fixture
def service(tmp_path):
    app = Service(tmp_path / "state")
    yield app
    app.close()


@pytest.fixture
def docx_file(tmp_path):
    path = tmp_path / "项目 说明.docx"
    document = Document()
    document.add_heading("原文标题", level=1)
    document.add_paragraph('独立原文 <script>alert("x")</script>')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "甲"
    table.cell(1, 1).text = "12"
    document.save(path)
    return path


def options(**overrides):
    return ConvertOptions(enable_ocr=False).with_overrides(**overrides).to_dict()


@pytest.mark.parametrize("folder", [False, True])
def test_native_import_dialog_filters_are_valid(service, folder):
    from webview.util import parse_file_type

    def dialog(dialog_type, **kwargs):
        for file_filter in kwargs["file_types"]:
            parse_file_type(file_filter)
        return None

    api = DesktopAPI(service)
    api._window = SimpleNamespace(create_file_dialog=dialog)
    assert api.choose_files(folder) == []


def convert(service, paths, **kwargs):
    files = service.register_files(paths)
    task = service.start_job([f["id"] for f in files], options(), **kwargs)
    return service.wait(task["id"]), files


def test_real_pdf_preview_conversion_and_no_overwrite(service, tmp_path):
    source = Path(__file__).parent / "samples" / "simple_test.pdf"
    prior = tmp_path / "simple_test.md"
    prior.write_text("user result", encoding="utf-8")
    job, files = convert(service, [source], output_dir=str(tmp_path))
    assert job["state"] == "completed"
    item = job["items"][0]
    assert Path(item["output_path"]).name == "simple_test_1.md"
    assert prior.read_text() == "user result"
    result = service.get_result(item["id"])
    assert result["markdown"] == Path(item["output_path"]).read_text()
    assert "<table>" in result["html"]
    preview = service.preview_source(files[0]["id"])
    assert preview["kind"] == "PDF" and preview["total"] >= 1
    assert base64.b64decode(preview["data_url"].split(",", 1)[1]).startswith(b"\x89PNG")


def test_pdf_continuous_preview_has_stable_sizes_and_renders_only_requested_page(service, tmp_path, monkeypatch):
    import pdfplumber
    import pypdfium2

    path = tmp_path / "mixed-pages.pdf"
    with pypdfium2.PdfDocument.new() as pdf:
        for width, height, rotation, crop in [
            (600, 800, 0, None),
            (900, 500, 0, None),
            (600, 800, 90, None),
            (600, 800, 0, (40, 60, 540, 720)),
        ]:
            page = pdf.new_page(width, height)
            if crop:
                page.set_cropbox(*crop)
            page.set_rotation(rotation)
            page.close()
        pdf.save(path)
    source = service.register_files([path])[0]
    rendered = []
    to_image = pdfplumber.page.Page.to_image

    def record_render(page, *args, **kwargs):
        rendered.append(page.page_number)
        return to_image(page, *args, **kwargs)

    monkeypatch.setattr(pdfplumber.page.Page, "to_image", record_render)
    first = service.preview_source(source["id"])
    assert first["total"] == 4
    assert len(first["page_sizes"]) == 4
    assert rendered == [1]  # Metadata must not eagerly rasterize the whole document.
    for number, size in enumerate(first["page_sizes"], start=1):
        preview = first if number == 1 else service.preview_source(source["id"], number)
        assert preview["page"] == number and preview["total"] == 4
        with Image.open(io.BytesIO(base64.b64decode(preview["data_url"].split(",", 1)[1]))) as image:
            assert size["width"] == pytest.approx(image.width, abs=2)
            assert size["height"] == pytest.approx(image.height, abs=2)
        if number > 1:
            assert "page_sizes" not in preview
    assert rendered == [1, 2, 3, 4]
    with pytest.raises(ValueError, match="页码超出"):
        service.preview_source(source["id"], 5)


def test_docx_original_is_read_independently(service, docx_file):
    job, files = convert(service, [docx_file])
    assert job["state"] == "completed"
    item = job["items"][0]
    assert item["pages"] is None
    original = service.preview_source(files[0]["id"])
    assert "原文标题" in original["html"] and "<table>" in original["html"]
    assert "<script>" not in original["html"]
    result = service.get_result(item["id"])
    document = Document(docx_file)
    document.add_paragraph("后来修改的原文")
    document.save(docx_file)
    updated = service.preview_source(files[0]["id"])
    assert "后来修改的原文" in updated["html"] and "改动" in updated["note"]
    assert "后来修改的原文" not in service.get_result(item["id"])["markdown"]
    assert "<script>" not in result["html"]


def test_history_survives_restart_and_missing_source(tmp_path, docx_file):
    app = Service(tmp_path / "history")
    job, files = convert(app, [docx_file])
    result_id = job["items"][0]["id"]
    expected = app.get_result(result_id)["markdown"]
    output = Path(job["items"][0]["output_path"])
    app.close()
    docx_file.rename(docx_file.with_name("moved.docx"))
    restarted = Service(tmp_path / "history")
    try:
        assert restarted.list_jobs()[0]["id"] == job["id"]
        assert restarted.get_result(result_id)["markdown"] == expected
        assert not restarted.get_result(result_id)["sources"][0]["exists"]
        with pytest.raises(ValueError, match="移动或删除"):
            restarted.preview_source(files[0]["id"])
        target = tmp_path / "export.html"
        restarted.export_to(result_id, target, "html")
        assert "原文标题" in target.read_text()
        restarted.delete_job(job["id"])
        assert restarted.list_jobs() == []
        assert output.exists() and target.exists()
        assert docx_file.with_name("moved.docx").exists()
    finally:
        restarted.close()


def test_word_all_text_table_is_present_in_export(service, tmp_path):
    path = tmp_path / "文字表格.docx"
    doc = Document()
    doc.add_heading("检查清单", level=1)
    table = doc.add_table(rows=3, cols=2)
    values_by_row = [["项目", "状态"], ["原文内容", "保留"], ["任务历史", "本地保存"]]
    for row, values in zip(table.rows, values_by_row, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    doc.save(path)
    job, _ = convert(service, [path])
    assert job["state"] == "completed"
    result = service.get_result(job["items"][0]["id"])
    assert "| 原文内容 | 保留 |" in result["markdown"]
    assert "| 任务历史 | 本地保存 |" in result["markdown"]
    assert "<table>" in result["html"]


def test_merge_order_and_partial_failure(service, tmp_path, docx_file):
    second = tmp_path / "另一份.docx"
    doc = Document()
    doc.add_paragraph("第二份正文")
    doc.save(second)
    bad = tmp_path / "损坏.docx"
    bad.write_bytes(b"not a document")
    job, _ = convert(service, [second, bad, docx_file], mode="merge", title="合并结果")
    assert job["state"] == "partial"
    item = job["items"][0]
    result = service.get_result(item["id"])
    assert "损坏.docx" in item["error"]
    assert result["markdown"].index("第二份正文") < result["markdown"].index("独立原文")
    assert len(result["sources"]) == 3 and item["pages"] is None


def test_stop_finishes_current_file(service, tmp_path, docx_file, monkeypatch):
    second = tmp_path / "second.docx"
    second.write_bytes(docx_file.read_bytes())
    entered, release = threading.Event(), threading.Event()
    parse = service._parse

    def blocked(source, opts):
        entered.set()
        assert release.wait(5)
        return parse(source, opts)

    monkeypatch.setattr(service, "_parse", blocked)
    files = service.register_files([docx_file, second])
    task = service.start_job([f["id"] for f in files], options())
    assert entered.wait(5)
    try:
        with pytest.raises(ValueError, match="处理中"):
            service.delete_job(task["id"])
        service.stop_job(task["id"])
    finally:
        release.set()
    job = service.wait(task["id"])
    assert job["state"] == "stopped"
    assert [item["state"] for item in job["items"]] == ["completed", "stopped"]
    assert Path(job["items"][0]["output_path"]).exists()
    assert not second.with_suffix(".md").exists()


def test_settings_and_task_snapshot_are_independent(service, docx_file):
    original = service.bootstrap()["settings"]
    files = service.register_files([docx_file])
    opts = options(remove_cjk_spaces=False)
    task = service.start_job([files[0]["id"]], opts)
    opts["remove_cjk_spaces"] = True
    changed = service.bootstrap()["settings"]
    changed["options"]["extract_images"] = True
    service.save_settings(changed)
    job = service.wait(task["id"])
    assert job["options"]["remove_cjk_spaces"] is False
    assert job["options"]["extract_images"] is False
    assert original["options"]["extract_images"] is False
    assert service.bootstrap()["settings"]["options"]["extract_images"] is True


def test_images_survive_export_and_output_removal(service, docx_file, tmp_path):
    image_path = tmp_path / "source.png"
    Image.new("RGB", (12, 10), "red").save(image_path)
    doc = Document(docx_file)
    doc.add_picture(str(image_path))
    doc.save(docx_file)
    source = service.register_files([docx_file])[0]
    job_id = service.start_job([source["id"]], options(extract_images=True))["id"]
    item = service.wait(job_id)["items"][0]
    result = service.get_result(item["id"])
    assert "data:image/png;base64" in result["html"]
    assert "![](assets_" in result["markdown"]
    exported = tmp_path / "exported" / "renamed.md"
    service.export_to(item["id"], exported)
    reference = exported.read_text().split("![](", 1)[1].split(")", 1)[0]
    assert (exported.parent / reference).is_file()
    Path(item["output_path"]).unlink()
    assert "data:image/png;base64" in service.get_result(item["id"])["html"]


@pytest.mark.parametrize("value", [[0, 1], [2, 1], [True, 3], [1], "1-2", [-1, 5]])
def test_invalid_page_ranges_rejected(value):
    with pytest.raises(ValueError):
        validate_options({"page_range": value})


def test_document_cannot_inject_script_or_fetch_remote_images():
    rendered = markdown_html(
        '<script>window.pywebview.api.copy_text("bad")</script>\n\n'
        "![](https://example.com/track.png)\n\n[link](https://example.com)"
    )
    assert "<script>" not in rendered and "<img" not in rendered and "href=" not in rendered


def test_interrupted_jobs_recovered(tmp_path):
    store = Store(tmp_path)
    store.create_job({"id": "j", "created_at": "2026-09-28", "state": "running"}, [{"id": "i", "state": "running"}])
    store.close()
    app = Service(tmp_path)
    try:
        job = app.list_jobs()[0]
        assert job["state"] == "interrupted" and job["items"][0]["state"] == "interrupted"
    finally:
        app.close()


def test_second_process_cannot_recover_live_history(tmp_path):
    lock = InstanceLock(tmp_path)
    try:
        with pytest.raises(RuntimeError, match="已在运行"):
            InstanceLock(tmp_path)
    finally:
        lock.close()
