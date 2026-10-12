"""Check the frozen bundle's real dependencies without using personal documents."""

import json
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

from desktop.version import get_version


def run(report_path):
    report_path = Path(report_path).resolve()
    report = {"ok": False, "platform": sys.platform, "checks": [], "ui_tested": False}
    try:
        report["version"] = get_version()
        from docx import Document

        from config import ConvertOptions
        from desktop.service import Service

        for asset in ("index.html", "app.css", "app.js", "assets/flipmd-symbol-v3-blue.png"):
            if not (Path(__file__).parent / "web" / asset).is_file():
                raise RuntimeError(f"缺少界面资源：{asset}")
        report["checks"].append("UI resources")

        if sys.platform == "win32":
            from desktop.windows import require_webview2

            require_webview2()
            from webview.guilib import initialize

            backend = initialize("edgechromium")
            if backend.renderer != "edgechromium":
                raise RuntimeError("没有加载 WebView2 引擎。")
            report["checks"].append("WebView2 / pythonnet assemblies")

        fixture = Path(__file__).parent / "samples" / "simple_test.pdf"
        if not fixture.exists():
            fixture = Path(__file__).parent.parent / "tests" / "samples" / "simple_test.pdf"
        with tempfile.TemporaryDirectory(prefix="flipmd-self-test-") as folder:
            root = Path(folder)
            pdf = root / "测试文档.pdf"
            shutil.copyfile(fixture, pdf)
            word = root / "测试文档.docx"
            document = Document()
            document.add_heading("Windows bundle check", level=1)
            table = document.add_table(rows=2, cols=2)
            for row, values in zip(table.rows, [["项目", "状态"], ["原文", "保留"]], strict=True):
                for cell, value in zip(row.cells, values, strict=True):
                    cell.text = value
            document.save(word)
            service = Service(root / "history")
            try:
                sources = service.register_files([pdf, word])
                task = service.start_job(
                    [source["id"] for source in sources],
                    ConvertOptions(enable_ocr=False).to_dict(),
                    output_dir=str(root / "output"),
                )
                job = service.wait(task["id"])
                if job["state"] != "completed":
                    raise RuntimeError(f"文档转换失败：{job}")
                result = service.get_result(job["items"][1]["id"])
                if "原文" not in result["markdown"] or "<table>" not in result["html"]:
                    raise RuntimeError("Word 表格未保留。")
                if not service.preview_source(sources[0]["id"])["data_url"].startswith("data:image/png"):
                    raise RuntimeError("PDF 渲染失败。")
                if "<table>" not in service.preview_source(sources[1]["id"])["html"]:
                    raise RuntimeError("Word 原文预览失败。")
                service.export_to(job["items"][1]["id"], root / "export.html", "html")
                report["checks"].extend(["PDF conversion / PDFium", "DOCX conversion / preview", "HTML export"])
            finally:
                service.close()
            reopened = Service(root / "history")
            try:
                if reopened.list_jobs()[0]["id"] != job["id"]:
                    raise RuntimeError("任务历史恢复失败。")
                report["checks"].append("SQLite history restart")
            finally:
                reopened.close()
        report["ok"] = True
    except Exception:
        report["error"] = traceback.format_exc()
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return 0 if report["ok"] else 1
