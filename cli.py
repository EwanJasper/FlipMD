"""命令行入口

用法：
    python cli.py convert <input> [-o output] [options]
    python cli.py batch <dir> [-o output_dir] [options]
    python cli.py info <input>

与 GUI 共享 engine.parse / word_engine.parse_docx / markdown_writer.render / config.ConvertOptions
"""

from __future__ import annotations

import argparse
import os
import sys

# Windows GBK 终端兼容：强制 stdout/stderr 用 UTF-8
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[union-attr]
    except Exception:
        pass

from config import ConvertOptions
from engine import get_page_count
from engine import parse as parse_pdf
from markdown_writer import render
from output_path import apply_name_template, resolve_unique_output_path
from word_engine import parse_docx

# ── 参数 → ConvertOptions ──────────────────────────────────────

def _build_options(args) -> ConvertOptions:
    """从命令行参数构造 ConvertOptions"""
    if args.no_postprocess:
        # 全关后处理
        opts = ConvertOptions()
        opts.remove_cjk_spaces = False
        opts.merge_broken_paragraphs = False
        opts.fix_headings = False
        opts.fix_list_number_positions = False
        opts.infer_empty_headers = False
        opts.merge_cross_page_tables = False
        opts.filter_footer = False
        opts.business_mode = False
    elif args.business_mode:
        opts = ConvertOptions.business()
    else:
        opts = ConvertOptions.default()

    # 覆盖单项开关
    if args.no_ocr:
        opts.enable_ocr = False
    if args.ocr_lang:
        opts.ocr_lang = args.ocr_lang
    if getattr(args, "extract_images", False):
        opts.extract_images = True
    if args.page_range:
        try:
            a, b = args.page_range.split("-")
            opts.page_range = (int(a.strip()), int(b.strip()))
        except (ValueError, AttributeError):
            print(f"⚠ 页码范围格式错误: {args.page_range}（应为 5-10），已忽略", file=sys.stderr)
    return opts


# ── 子命令实现 ─────────────────────────────────────────────────

def _convert_one(path: str, output: str | None, options: ConvertOptions,
                 fmt: str = "md") -> tuple[bool, str, str]:
    """转换单个文件。返回 (success, output_path, message)"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        try:
            page_count = get_page_count(path)
        except Exception:
            page_count = 0
        blocks = parse_pdf(path, options)
        page_info = f"{page_count} 页" if page_count else "-"
    elif ext == ".docx":
        page_count = 0
        blocks = parse_docx(path, options)
        page_info = "-"
    elif ext == ".doc":
        return False, "", "不支持 .doc 格式，请先用 Word 另存为 .docx"
    else:
        return False, "", f"不支持的文件格式: {ext}"

    # 默认输出路径：与源同目录、同名（.md 或 .html）
    out_ext = ".html" if fmt == "html" else ".md"
    if not output:
        base = os.path.splitext(path)[0]
        output = base + out_ext

    # 图片提取（导出到 output 同级 images/ 目录）
    img_count = 0
    if options.extract_images:
        try:
            from image_extract import extract_images
            img_blocks = extract_images(path, output)
            if img_blocks:
                # 合并到 blocks 中（按 y_top 排序保持阅读顺序）
                blocks = blocks + img_blocks
                blocks.sort(key=lambda b: b.y_top)
                img_count = len(img_blocks)
        except Exception as e:
            print(f"  ⚠ 图片提取失败: {e}", file=sys.stderr)

    title = os.path.splitext(os.path.basename(path))[0]
    if fmt == "html":
        from html_writer import render_html
        content = render_html(blocks, title=title)
    else:
        metadata = None
        if options.add_front_matter:
            from datetime import date
            metadata = {
                "source": os.path.basename(path),
                "format": ext.lstrip(".").upper(),
                "pages": page_count,
                "converted": date.today().isoformat(),
            }
        content = render(blocks, title=title, options=options, metadata=metadata)

    os.makedirs(os.path.dirname(output) or ".", exist_ok=True)
    with open(output, "w", encoding="utf-8") as f:
        f.write(content)

    tables = sum(1 for b in blocks if b.type == "table")
    paras = sum(1 for b in blocks if b.type == "paragraph")
    lines = len(content.splitlines())
    img_part = f" | {img_count} 图" if img_count else ""
    msg = f"  ✓ {page_info} | {paras} 段 | {tables} 表{img_part} | {lines} 行 → {output}"
    return True, output, msg


def cmd_convert(args) -> int:
    """转换单个文件"""
    if not os.path.exists(args.input):
        print(f"✗ 文件不存在: {args.input}", file=sys.stderr)
        return 2
    options = _build_options(args)
    print(f"[转换] {args.input}")
    ok, out, msg = _convert_one(args.input, args.output, options, fmt=args.format)
    if ok:
        print(msg)
        return 0
    print(f"  ✗ 失败: {msg}", file=sys.stderr)
    return 1


def cmd_batch(args) -> int:
    """批量转换目录"""
    if not os.path.isdir(args.directory):
        print(f"✗ 目录不存在: {args.directory}", file=sys.stderr)
        return 2

    # 收集文件
    inputs: list[str] = []
    if args.recursive:
        for root, _dirs, files in os.walk(args.directory):
            for fname in files:
                if fname.lower().endswith((".pdf", ".docx")):
                    inputs.append(os.path.join(root, fname))
    else:
        for fname in sorted(os.listdir(args.directory)):
            if fname.lower().endswith((".pdf", ".docx")):
                inputs.append(os.path.join(args.directory, fname))

    if not inputs:
        print(f"⚠ 目录中没有 PDF/DOCX 文件: {args.directory}")
        return 0

    options = _build_options(args)
    print(f"[批量] {args.directory}  共 {len(inputs)} 个文件")

    success = 0
    failed = 0
    out_ext = ".html" if args.format == "html" else ".md"
    for idx, path in enumerate(inputs, 1):
        # 计算输出路径（含重名保护）
        out_path = resolve_unique_output_path(
            args.output_dir, path, ext=out_ext, template=args.name_template
        )

        print(f"[{idx}/{len(inputs)}] {os.path.basename(path)}")
        try:
            ok, out, msg = _convert_one(path, out_path, options, fmt=args.format)
            if ok:
                print(msg)
                success += 1
            else:
                print(f"  ✗ 失败: {msg}", file=sys.stderr)
                failed += 1
        except Exception as e:
            print(f"  ✗ 异常: {e}", file=sys.stderr)
            import traceback
            traceback.print_exc()
            failed += 1

    print("=" * 50)
    print(f"✓ 成功 {success} / 失败 {failed} / 共 {len(inputs)}")
    return 0 if failed == 0 else 1


def cmd_info(args) -> int:
    """显示文档信息"""
    path = args.input
    if not os.path.exists(path):
        print(f"✗ 文件不存在: {path}", file=sys.stderr)
        return 2

    ext = os.path.splitext(path)[1].lower()
    size = os.path.getsize(path)
    size_str = _format_size(size)

    print(f"文件: {path}")
    print(f"格式: {ext.upper().lstrip('.')}")
    print(f"大小: {size_str}")

    if ext == ".pdf":
        try:
            pages = get_page_count(path)
            print(f"页数: {pages}")
        except Exception as e:
            print(f"页数: 读取失败 ({e})")
        # 简单文字密度采样
        try:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                n = min(5, len(pdf.pages))
                total = 0
                for i in range(n):
                    words = pdf.pages[i].extract_words()
                    total += len(words)
                avg = total / n if n else 0
                print(f"前 {n} 页平均字数: {avg:.1f}")
                if avg < 10:
                    print("判定: 可能是扫描型 PDF（需要 OCR）")
                else:
                    print("判定: 文本型 PDF")
        except Exception:
            pass
    elif ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            paras = len(doc.paragraphs)
            tables = len(doc.tables)
            print(f"段落数: {paras}")
            print(f"表格数: {tables}")
        except Exception as e:
            print(f"读取失败: {e}")
    else:
        print(f"不支持的格式: {ext}")
        return 1

    return 0


# ── 辅助函数 ───────────────────────────────────────────────────

def _apply_name_template(template: str | None, base_name: str) -> str:
    """应用文件名模板（兼容旧测试入口，委派给 output_path 模块）"""
    return apply_name_template(template, base_name)


def _format_size(size_bytes: float) -> str:
    for unit in ["B", "KB", "MB"]:
        if size_bytes < 1024:
            return f"{size_bytes:.1f} {unit}"
        size_bytes /= 1024
    return f"{size_bytes:.1f} GB"


def _resolve_output_path(out_dir: str | None, src_path: str, name_template: str | None) -> str:
    """根据输出目录和模板计算输出路径（委派给 output_path 模块）"""
    return resolve_unique_output_path(out_dir, src_path, ext=".md", template=name_template)


# ── argparse ───────────────────────────────────────────────────

def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="pdf2md",
        description="PDF & Word → Markdown 转换工具（CLI）",
    )
    sub = p.add_subparsers(dest="command", required=True, metavar="<command>")

    # convert
    p_conv = sub.add_parser("convert", help="转换单个文件")
    p_conv.add_argument("input", help="输入文件路径（.pdf / .docx）")
    p_conv.add_argument("-o", "--output", help="输出 Markdown 文件路径（默认与源同目录）")
    _add_common_options(p_conv)
    p_conv.set_defaults(func=cmd_convert)

    # batch
    p_batch = sub.add_parser("batch", help="批量转换目录")
    p_batch.add_argument("directory", help="输入目录")
    p_batch.add_argument("-o", "--output-dir", help="输出目录（默认与源同目录）")
    p_batch.add_argument("--recursive", action="store_true", help="递归子目录")
    p_batch.add_argument("--name-template", help="输出文件名模板（{name}/{date}/{seq}）")
    _add_common_options(p_batch)
    p_batch.set_defaults(func=cmd_batch)

    # info
    p_info = sub.add_parser("info", help="显示文档信息")
    p_info.add_argument("input", help="输入文件路径")
    p_info.set_defaults(func=cmd_info)

    return p


def _add_common_options(p: argparse.ArgumentParser):
    p.add_argument("--no-postprocess", action="store_true",
                   help="禁用所有后处理")
    p.add_argument("--business-mode", action="store_true",
                   help="启用业务文档模式（社会保单等特定文档）")
    p.add_argument("--no-ocr", action="store_true",
                   help="禁用 OCR（扫描型 PDF 不再自动 OCR）")
    p.add_argument("--ocr-lang", choices=["ch", "en", "ch_en"],
                   help="OCR 语言")
    p.add_argument("--extract-images", action="store_true",
                   help="提取图片到 output/images/ 并在 Markdown 中引用")
    p.add_argument("--format", choices=["md", "html"], default="md",
                   help="输出格式（默认 md）")
    p.add_argument("--page-range", metavar="START-END",
                   help="页码范围，如 5-10")


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
